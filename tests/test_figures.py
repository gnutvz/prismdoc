"""Tests for T-011 figure sub-pipeline (extract -> process -> merge)."""

from __future__ import annotations

import base64
import builtins
from pathlib import Path

import fitz
import pytest
from openpyxl import Workbook
from PIL import Image

import prismdoc
from prismdoc import (
    Context,
    Document,
    Figure,
    FigureExtractStage,
    FigureMergeStage,
    FigureProcessStage,
    FigureProcessor,
    IngestStage,
    ParseStage,
    Pipeline,
    Source,
    registry,
)
from prismdoc.stages.figures import (
    OcrFigureProcessor,
    StubFigureProcessor,
    register_plugins,
)


def _make_png(path: Path, size: tuple[int, int] = (32, 24)) -> None:
    Image.new("RGB", size, color=(0, 128, 255)).save(path)


def _make_pdf_with_embedded_image(path: Path, png_path: Path) -> None:
    doc = fitz.open()
    page = doc.new_page()
    page.insert_text((72, 72), "Catalog page with a figure")
    page.insert_image(fitz.Rect(100, 100, 200, 180), filename=str(png_path))
    doc.save(path)
    doc.close()


def _make_text_only_pdf(path: Path) -> None:
    doc = fitz.open()
    page = doc.new_page()
    page.insert_text((72, 72), "No figures here")
    doc.save(path)
    doc.close()


def _make_xlsx(path: Path) -> None:
    wb = Workbook()
    sheet = wb.active
    assert sheet is not None
    sheet.append(["sku", "price"])
    sheet.append(["A1", "9.99"])
    wb.save(path)


def test_figure_extract_inserts_placeholder_and_stores_figure(
    tmp_path: Path,
) -> None:
    png_path = tmp_path / "fig.png"
    pdf_path = tmp_path / "with_fig.pdf"
    _make_png(png_path)
    _make_pdf_with_embedded_image(pdf_path, png_path)

    doc = Document(source=Source(path=str(pdf_path), mime="application/pdf"))
    doc = Pipeline([IngestStage(), ParseStage()]).run(doc, Context())
    before = doc.artifacts["parsed_markdown"]
    assert "[[FIGURE:" not in before
    assert "Catalog page with a figure" in before

    result = FigureExtractStage().run(doc, Context())
    md = result.artifacts["parsed_markdown"]
    assert md.count("[[FIGURE:fig_0_0]]") == 1
    assert "[[FIGURE:" in md
    # Placeholder sits after page content (same page section).
    assert md.index("Catalog page with a figure") < md.index("[[FIGURE:fig_0_0]]")

    figures = result.artifacts["figures"]
    assert len(figures) == 1
    figure = Figure.model_validate(figures[0])
    assert figure.id == "fig_0_0"
    assert figure.page_index == 0
    assert figure.width == 32
    assert figure.height == 24
    assert figure.mime.startswith("image/")
    assert base64.b64decode(figure.image_b64)
    assert figure.result is None


def test_figure_process_stub_sets_result(tmp_path: Path) -> None:
    png_path = tmp_path / "fig.png"
    pdf_path = tmp_path / "with_fig.pdf"
    _make_png(png_path)
    _make_pdf_with_embedded_image(pdf_path, png_path)

    doc = Document(source=Source(path=str(pdf_path), mime="application/pdf"))
    doc = Pipeline(
        [IngestStage(), ParseStage(), FigureExtractStage()]
    ).run(doc, Context())

    result = FigureProcessStage().run(doc, Context())
    figure = Figure.model_validate(result.artifacts["figures"][0])
    assert figure.result == "[figure fig_0_0: 32x24 image/png]"


def test_figure_merge_replaces_placeholder(tmp_path: Path) -> None:
    png_path = tmp_path / "fig.png"
    pdf_path = tmp_path / "with_fig.pdf"
    _make_png(png_path)
    _make_pdf_with_embedded_image(pdf_path, png_path)

    doc = Document(source=Source(path=str(pdf_path), mime="application/pdf"))
    doc = Pipeline(
        [
            IngestStage(),
            ParseStage(),
            FigureExtractStage(),
            FigureProcessStage(),
        ]
    ).run(doc, Context())
    assert "[[FIGURE:fig_0_0]]" in doc.artifacts["parsed_markdown"]

    result = FigureMergeStage().run(doc, Context())
    md = result.artifacts["parsed_markdown"]
    assert "[[FIGURE:" not in md
    assert "[figure fig_0_0: 32x24 image/png]" in md
    assert "Catalog page with a figure" in md


def test_full_figure_pipeline_round_trip(tmp_path: Path) -> None:
    png_path = tmp_path / "fig.png"
    pdf_path = tmp_path / "with_fig.pdf"
    _make_png(png_path, size=(40, 30))
    _make_pdf_with_embedded_image(pdf_path, png_path)

    doc = Document(source=Source(path=str(pdf_path), mime="application/pdf"))
    after_parse = Pipeline([IngestStage(), ParseStage()]).run(
        doc.model_copy(deep=True), Context()
    )
    assert "[[FIGURE:" not in after_parse.artifacts["parsed_markdown"]

    after_extract = FigureExtractStage().run(after_parse, Context())
    assert "[[FIGURE:fig_0_0]]" in after_extract.artifacts["parsed_markdown"]

    result = Pipeline(
        [FigureProcessStage(), FigureMergeStage()]
    ).run(after_extract, Context())
    md = result.artifacts["parsed_markdown"]
    assert "[[FIGURE:" not in md
    assert "[figure fig_0_0: 40x30 image/png]" in md
    # Result appears in the page-0 section (after original text).
    assert md.index("Catalog page with a figure") < md.index(
        "[figure fig_0_0: 40x30 image/png]"
    )


def test_figure_less_text_pdf_passthrough(tmp_path: Path) -> None:
    pdf_path = tmp_path / "plain.pdf"
    _make_text_only_pdf(pdf_path)

    doc = Document(source=Source(path=str(pdf_path), mime="application/pdf"))
    result = Pipeline(
        [
            IngestStage(),
            ParseStage(),
            FigureExtractStage(),
            FigureProcessStage(),
            FigureMergeStage(),
        ]
    ).run(doc, Context())

    assert result.artifacts.get("figures") == []
    assert "[[FIGURE:" not in result.artifacts["parsed_markdown"]
    assert "No figures here" in result.artifacts["parsed_markdown"]


def test_figure_less_xlsx_passthrough(tmp_path: Path) -> None:
    xlsx_path = tmp_path / "sheet.xlsx"
    _make_xlsx(xlsx_path)

    doc = Document(
        source=Source(
            path=str(xlsx_path),
            mime=(
                "application/vnd.openxmlformats-officedocument"
                ".spreadsheetml.sheet"
            ),
        )
    )
    before = Pipeline([IngestStage(), ParseStage()]).run(doc, Context())
    md_before = before.artifacts["parsed_markdown"]

    result = Pipeline(
        [FigureExtractStage(), FigureProcessStage(), FigureMergeStage()]
    ).run(before, Context())

    assert result.artifacts.get("figures") == []
    assert result.artifacts["parsed_markdown"] == md_before
    assert "[[FIGURE:" not in result.artifacts["parsed_markdown"]


def test_merge_unprocessed_fallback() -> None:
    doc = Document(
        source=Source(path="/tmp/x.pdf", mime="application/pdf"),
        artifacts={
            "parsed_markdown": "before\n[[FIGURE:fig_0_0]]\nafter",
            "figures": [
                Figure(
                    id="fig_0_0",
                    page_index=0,
                    width=1,
                    height=1,
                    image_b64=base64.b64encode(b"x").decode("ascii"),
                    result=None,
                ).model_dump()
            ],
        },
    )
    result = FigureMergeStage().run(doc, Context())
    assert "[[FIGURE:" not in result.artifacts["parsed_markdown"]
    assert "[unprocessed figure fig_0_0]" in result.artifacts["parsed_markdown"]


def test_ocr_processor_raises_clear_import_error(
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    real_import = builtins.__import__

    def fake_import(name: str, *args: object, **kwargs: object) -> object:
        if name == "rapidocr_onnxruntime" or name.startswith(
            "rapidocr_onnxruntime."
        ):
            raise ImportError("No module named 'rapidocr_onnxruntime'")
        return real_import(name, *args, **kwargs)

    monkeypatch.setattr(builtins, "__import__", fake_import)

    figure = Figure(
        id="fig_0_0",
        page_index=0,
        width=1,
        height=1,
        image_b64=base64.b64encode(b"x").decode("ascii"),
    )
    with pytest.raises(ImportError, match=r"prismdoc\[docling\]"):
        OcrFigureProcessor().process(figure)


def test_figure_exports_and_registry() -> None:
    assert issubclass(FigureExtractStage, object)
    assert issubclass(FigureProcessStage, object)
    assert issubclass(FigureMergeStage, object)
    assert issubclass(FigureProcessor, object)
    assert issubclass(StubFigureProcessor, FigureProcessor)
    assert callable(prismdoc.Figure)
    assert callable(prismdoc.FigureExtractStage)
    assert callable(prismdoc.FigureProcessStage)
    assert callable(prismdoc.FigureMergeStage)
    assert callable(prismdoc.FigureProcessor)

    register_plugins()
    keys = registry.get_keys()
    assert "figures.extract" in keys
    assert "figures.process" in keys
    assert "figures.merge" in keys

    assert isinstance(registry.create("figures.extract"), FigureExtractStage)
    assert isinstance(registry.create("figures.process"), FigureProcessStage)
    assert isinstance(registry.create("figures.merge"), FigureMergeStage)
