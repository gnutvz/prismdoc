"""Tests for T-002 ingest loaders (PDF, image, xlsx) and IngestStage."""

from __future__ import annotations

from pathlib import Path

import fitz
import pytest
from openpyxl import Workbook
from PIL import Image

from prismdoc import Context, Document, IngestStage, Loader, Pipeline, Source, registry
from prismdoc.stages.ingest import ImageLoader, PdfLoader, XlsxLoader, register_plugins


def _make_pdf(path: Path, pages: list[str]) -> None:
    doc = fitz.open()
    for text in pages:
        page = doc.new_page()
        page.insert_text((72, 72), text)
    doc.save(path)
    doc.close()


def _make_png(path: Path) -> None:
    Image.new("RGB", (8, 8), color=(255, 0, 0)).save(path)


def _make_xlsx(path: Path) -> None:
    wb = Workbook()
    sheet1 = wb.active
    assert sheet1 is not None
    sheet1.title = "Prices"
    sheet1.append(["sku", "price"])
    sheet1.append(["A1", "9.99"])
    sheet2 = wb.create_sheet("Stock")
    sheet2.append(["sku", "qty"])
    sheet2.append(["A1", "3"])
    wb.save(path)


def test_pdf_loader_two_pages(tmp_path: Path) -> None:
    pdf_path = tmp_path / "sample.pdf"
    _make_pdf(pdf_path, ["Hello PDF page 1", "Hello PDF page 2"])

    pages = PdfLoader().load(Source(path=str(pdf_path), mime="application/pdf"))

    assert len(pages) == 2
    assert pages[0].index == 0
    assert pages[1].index == 1
    assert "Hello PDF page 1" in pages[0].text
    assert "Hello PDF page 2" in pages[1].text
    assert pages[0].blocks
    assert pages[0].blocks[0].bbox is not None
    assert len(pages[0].blocks[0].bbox) == 4


def test_image_loader_attaches_ref_without_ocr(tmp_path: Path) -> None:
    img_path = tmp_path / "sample.png"
    _make_png(img_path)

    pages = ImageLoader().load(Source(path=str(img_path), mime="image/png"))

    assert len(pages) == 1
    assert pages[0].index == 0
    assert pages[0].text == ""
    assert pages[0].image_ref == str(img_path)


def test_xlsx_loader_two_sheets(tmp_path: Path) -> None:
    xlsx_path = tmp_path / "sample.xlsx"
    _make_xlsx(xlsx_path)

    pages = XlsxLoader().load(
        Source(
            path=str(xlsx_path),
            mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        )
    )

    assert len(pages) == 2
    assert pages[0].index == 0
    assert pages[1].index == 1
    assert "sku\tprice" in pages[0].text
    assert "A1\t9.99" in pages[0].text
    assert "sku\tqty" in pages[1].text
    assert "A1\t3" in pages[1].text
    assert pages[0].blocks[0].text == "Prices"
    assert pages[0].blocks[0].meta.get("kind") == "sheet_name"
    assert pages[1].blocks[0].text == "Stock"


def test_ingest_stage_via_pipeline(tmp_path: Path) -> None:
    pdf_path = tmp_path / "pipeline.pdf"
    _make_pdf(pdf_path, ["Pipeline page"])

    doc = Document(source=Source(path=str(pdf_path), mime="application/pdf"))
    result = Pipeline([IngestStage()]).run(doc, Context())

    assert len(result.pages) == 1
    assert "Pipeline page" in result.pages[0].text
    assert len(result.trace) == 1
    assert result.trace[0].stage == "ingest"
    assert result.trace[0].ok is True
    assert result.trace[0].duration_ms >= 0


def test_ingest_unsupported_extension(tmp_path: Path) -> None:
    bad_path = tmp_path / "notes.txt"
    bad_path.write_text("not supported", encoding="utf-8")
    doc = Document(source=Source(path=str(bad_path), mime="text/plain"))

    with pytest.raises(ValueError, match="Unsupported source extension"):
        IngestStage().run(doc, Context())


def test_ingest_stage_raises_input_too_large_when_max_pages_exceeded(
    tmp_path: Path,
) -> None:
    from prismdoc import InputTooLargeError

    pdf_path = tmp_path / "many.pdf"
    _make_pdf(pdf_path, ["page one", "page two", "page three"])

    doc = Document(source=Source(path=str(pdf_path), mime="application/pdf"))
    ctx = Context(options={"max_pages": 2})

    with pytest.raises(InputTooLargeError, match="3 pages"):
        IngestStage().run(doc, ctx)

    assert len(doc.pages) == 3


def test_ingest_stage_respects_max_pages_when_within_limit(
    tmp_path: Path,
) -> None:
    pdf_path = tmp_path / "ok.pdf"
    _make_pdf(pdf_path, ["only one"])

    doc = Document(source=Source(path=str(pdf_path), mime="application/pdf"))
    result = IngestStage().run(doc, Context(options={"max_pages": 2}))

    assert len(result.pages) == 1


def test_ingest_exports_and_registry() -> None:
    assert issubclass(IngestStage, object)
    assert issubclass(Loader, object)

    register_plugins()
    keys = registry.get_keys()
    assert "loader.pdf" in keys
    assert "loader.image" in keys
    assert "loader.xlsx" in keys
    assert "ingest.default" in keys

    stage = registry.create("ingest.default")
    assert isinstance(stage, IngestStage)
    assert isinstance(registry.create("loader.pdf"), PdfLoader)
